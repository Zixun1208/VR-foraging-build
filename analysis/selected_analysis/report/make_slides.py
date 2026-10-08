#!/usr/bin/env python3
"""Slides (.pptx) from ``runs/summary``; the content is in deck_spec.py.

    python -m selected_analysis.report.make_slides --dir runs/summary        -> runs/summary/foraging_summary.pptx

Native charts (python-pptx), 16:9 at 10 x 5.625 in. Run summary_stats.py and summary_md.py first
(summary_md.py copies the figures the slides use into ``<dir>/figs``).
"""
from __future__ import annotations

import argparse
import io
import json
import os

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .. import config
from . import deck_spec
rgb = lambda h: RGBColor.from_string(h)
INK, TEAL, ORANGE = rgb(deck_spec.INK), rgb(deck_spec.TEAL), rgb(deck_spec.ORANGE)
MUTED, GRID, CARD, WHITE = rgb("52666B"), rgb("DFE6E4"), rgb("F3F7F6"), rgb("FFFFFF")
HEAD, BODY = "Cambria", "Calibri"


def run_text(tf, parts, size=14, color=INK, bold=False, font=BODY, align=None):
    """parts: str or list of (text, {bold,color,size}) paragraphs."""
    if isinstance(parts, str):
        parts = [(parts, {})]
    for i, (text, o) in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = text
        f = r.font
        f.name, f.size = o.get("font", font), Pt(o.get("size", size))
        f.bold, f.color.rgb = o.get("bold", bold), o.get("color", color)
        if align:
            p.alignment = align
        p.space_after = Pt(6)


def text(slide, x, y, w, h, parts, anchor=MSO_ANCHOR.TOP, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    run_text(tf, parts, **kw)
    return tb


def title(slide, t):
    text(slide, 0.5, 0.35, 9, 0.8, t, size=28, color=TEAL, bold=True, font=HEAD)


def card(slide, x, y, w, h):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = CARD
    s.line.fill.background()
    return s


def picture(slide, path, x, y, box_w, box_h, crop=None, alt=""):
    im = Image.open(path)
    if crop:
        im = im.crop((0, int(im.height * crop[0]), im.width, int(im.height * crop[1])))
    k = min(box_w / im.width, box_h / im.height)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    pic = slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(im.width * k), Inches(im.height * k))
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def chart(slide, kind, spec, x, y, w, h):
    kinds = {"col": XL_CHART_TYPE.COLUMN_CLUSTERED, "stacked": XL_CHART_TYPE.COLUMN_STACKED,
             "bar_h": XL_CHART_TYPE.BAR_CLUSTERED}
    cd = CategoryChartData()
    cd.categories = spec["cats"]
    for name, vals in spec["series"]:
        cd.add_series(name, vals)
    ch = slide.shapes.add_chart(kinds[kind], Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.font.name, ch.font.size, ch.font.color.rgb = BODY, Pt(12), MUTED
    multi = len(spec["series"]) > 1
    ch.has_legend = multi
    if multi:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size, ch.legend.font.color.rgb = Pt(12), INK
    ch.has_title = bool(spec.get("title"))
    if spec.get("title"):
        ch.chart_title.text_frame.text = spec["title"]
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(13), False, INK, BODY
    for s, c in zip(ch.plots[0].series, spec["colors"]):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(11)
    if kind == "stacked":
        dl.position, dl.font.color.rgb = XL_LABEL_POSITION.CENTER, WHITE
    else:
        dl.position, dl.font.color.rgb = XL_LABEL_POSITION.OUTSIDE_END, INK
    va, ca = ch.value_axis, ch.category_axis
    va.major_gridlines.format.line.color.rgb = GRID
    va.format.line.fill.background()
    ca.format.line.color.rgb = GRID
    ca.has_major_gridlines = False
    va.tick_labels.font.size = ca.tick_labels.font.size = Pt(12)
    if spec.get("ymax"):
        va.maximum_scale, va.minimum_scale = spec["ymax"], 0
    if spec.get("ylabel"):
        va.has_title = True
        va.axis_title.text_frame.text = spec["ylabel"]
        r = va.axis_title.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb = Pt(12), False, MUTED
    if kind == "bar_h":
        ca.reverse_order = True
    return ch


def render(prs, sp):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK if sp["type"] == "title" else WHITE
    t = sp["type"]
    if t == "title":
        text(s, 0.7, 1.5, 8.6, 1.6, sp["title"], size=40, color=WHITE, bold=True, font=HEAD)
        text(s, 0.7, 3.3, 8.0, 0.9, sp["sub"], size=18, color=rgb("CADCDD"))
        return s
    title(s, sp["title"])
    if sp.get("subtitle"):
        text(s, 0.5, 0.88, 9, 0.3, sp["subtitle"], size=14, color=MUTED)
    if t == "bar_h":
        chart(s, "bar_h", sp["chart"], 0.5, 1.2, 5.9, 4.0)
        text(s, 6.8, 1.3, 2.7, 3.8, [(a, {"bold": b}) for a, b in sp["side"]])
    elif t == "cards":
        for i, (n, head, body) in enumerate(sp["cards"]):
            x, y = 0.5 + (i % 2) * 4.6, 1.3 + (i // 2) * 1.95
            card(s, x, y, 4.3, 1.7)
            b = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(y + 0.2), Inches(0.5), Inches(0.5))
            b.fill.solid()
            b.fill.fore_color.rgb = ORANGE
            b.line.fill.background()
            text(s, x + 0.2, y + 0.2, 0.5, 0.5, n, size=16, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
            text(s, x + 0.85, y + 0.2, 3.3, 0.5, head, size=18, color=TEAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + 0.2, y + 0.85, 3.9, 0.8, body)
    elif t == "image_stack":
        for j, (path, label, crop) in enumerate(sp["images"]):
            y = 1.15 + j * 2.05
            text(s, 3.3, y, 5.9, 0.3, label, size=12, color=TEAL, bold=True)
            picture(s, path, 3.3, y + 0.3, 5.9, 1.75, crop=crop, alt=label)
        text(s, 0.5, 1.4, 2.6, 3.4, [(x, {}) for x in sp["side"]])
    elif t == "chart":
        chart(s, sp["kind"], sp["chart"], 0.5, 1.2, 9, 3.4)
        text(s, 0.5, 4.75, 9, 0.7, sp["caption"])
    elif t == "chart_image":
        chart(s, "stacked", sp["chart"], 0.5, 1.2, 5.4, 3.9)
        picture(s, sp["image"], 6.1, 1.3, 3.4, 2.7, alt="Difference between the two patches, 20/100 task")
        text(s, 6.1, 4.1, 3.4, 1.2, sp["note"], size=12)
    elif t == "chart_stats":
        chart(s, sp["kind"], sp["chart"], 0.5, 1.2, 5.9, 3.9)
        for i, (big, label, col) in enumerate(sp["stats"]):
            text(s, 6.8, 1.3 + i * 1.8, 2.7, 0.9, big, size=48, color=rgb(col), bold=True)
            text(s, 6.8, 2.2 + i * 1.8, 2.7, 0.8, label)
    elif t == "image_stats":
        picture(s, sp["image"], 0.5, 1.15, 9, 2.7, alt=sp["title"])
        for i, (big, label, col) in enumerate(sp["stats"]):
            x = 0.5 + i * 3.1
            text(s, x, 4.0, 2.9, 0.7, big, size=36, color=rgb(col), bold=True)
            text(s, x, 4.7, 2.9, 0.7, label)
    elif t == "image_grid":
        for c in sp["cells"]:
            text(s, c["lx"], c["ly"], c["lw"], 0.3, c["label"], size=12, color=TEAL, bold=True)
            picture(s, c["path"], c["ix"], c["iy"], c["iw"], c["ih"], crop=c["crop"], alt=c["label"])
        if sp.get("caption"):
            text(s, 0.5, 4.85, 9.0, 0.6, sp["caption"], size=12)
    elif t == "stats":
        for i, (big, label, col) in enumerate(sp["stats"]):
            x = 0.5 + i * 3.1
            text(s, x, 1.9, 2.9, 1.0, big, size=32, color=rgb(col), bold=True)
            text(s, x, 3.0, 2.9, 0.9, label)
        text(s, 0.5, 4.5, 9.0, 0.8, sp["note"], size=13)
    elif t == "image_text" and sp.get("layout") == "below":
        picture(s, sp["image"], 0.5, 1.1, 9.0, 3.1, alt=sp["title"])
        text(s, 0.5, 4.3, 9.0, 1.2, [(x, {}) for x in sp["text"]], size=13)
    elif t == "image_text":
        picture(s, sp["image"], 0.5, 1.2, 6.0, 4.0, alt=sp["title"])
        text(s, 6.8, 1.4, 2.7, 3.8, [(x, {}) for x in sp["text"]], size=14)
    elif t == "columns":
        for i, (head, a1, a2, col) in enumerate(sp["cols"]):
            x = 0.5 + i * 3.1
            card(s, x, 1.3, 2.9, 3.1)
            text(s, x + 0.2, 1.45, 2.5, 0.5, head, size=20, color=rgb(col), bold=True, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + 0.2, 2.05, 2.5, 0.9, a1)
            text(s, x + 0.2, 3.0, 2.5, 1.3, a2, bold=True)
        text(s, 0.5, 4.65, 9, 0.5, sp["footer"])
    elif t == "bullets":
        for x, (head, col, items) in ((0.5, sp["left"]), (5.2, sp["right"])):
            tb = s.shapes.add_textbox(Inches(x), Inches(1.3), Inches(4.3), Inches(3.9))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            run_text(tf, [(head, {"bold": True, "color": rgb(col)})])
            for it in items:
                p = tf.add_paragraph()
                r = p.add_run()
                r.text = "•  " + it
                r.font.name, r.font.size, r.font.color.rgb = BODY, Pt(14), INK
                p.space_after = Pt(6)
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=config.SUMMARY)
    a = ap.parse_args()
    with open(os.path.join(a.dir, "summary.json")) as f:
        summ = json.load(f)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    prs.core_properties.title = "Patch leaving across all foraging datasets"
    for sp in deck_spec.slides(summ, os.path.join(a.dir, "figs")):
        render(prs, sp)
    out = os.path.join(a.dir, "foraging_summary.pptx")
    prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
